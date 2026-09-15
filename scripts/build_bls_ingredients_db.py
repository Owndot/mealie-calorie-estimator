#!/usr/bin/env python3
"""Build a compact, lossless SQLite subset of BLS 4.0 for recipe ingredients.

The selected food codes are the canonical BLS codes from EiSiMo/bls-icons
(items.csv), restricted to BLS main groups B-W. X/Y menu components are
excluded. The build is intentionally strict and expects exactly 2,992 foods.

Nutrient values are copied from the official BLS 4.0 workbook without
normalisation. Numeric values are stored as their textual representation;
special BLS tokens such as TR, <LOD, <LOQ and '-' remain unchanged. Data
origin/reference columns are deliberately omitted.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import re
import sqlite3
import sys
import tempfile
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any

import requests
from openpyxl import load_workbook

BLS_HOME = "https://blsdb.de/"
BLS_DOWNLOAD = "https://blsdb.de/download"
CANONICAL_ITEMS_URL = "https://raw.githubusercontent.com/EiSiMo/bls-icons/master/items.csv"
EXPECTED_FOODS = 2992
EXPECTED_NUTRIENTS = 138
ALLOWED_MAIN_GROUPS = frozenset("BCDEFGHKMNPQRSTUVW")
ATTRIBUTION = (
    "Max Rubner-Institut (2025): Bundeslebensmittelschlüssel (BLS), "
    "Version 4.0 - Deutsche Nährstoffdatenbank. Karlsruhe. "
    "DOI: 10.25826/Data20251217-134202-0"
)


def download_official_xlsx() -> bytes:
    """Download the current official BLS 4.0 ZIP and return the data workbook."""
    session = requests.Session()
    session.headers.update({"User-Agent": "Mozilla/5.0 (BLS ingredient DB builder)"})

    # Establish the session/cookie first, then obtain the rotating token/link.
    session.get(BLS_HOME, timeout=60).raise_for_status()
    page = session.get(BLS_DOWNLOAD, timeout=60)
    page.raise_for_status()

    match = re.search(
        r'href=["\']([^"\']*BLS[^"\']*\.zip[^"\']*)["\']',
        page.text,
        flags=re.IGNORECASE,
    )
    if not match:
        # The page may expose the token separately and construct the stable ZIP path.
        token_match = re.search(r"[?&]token=([A-Za-z0-9-]{8,})", page.text)
        if not token_match:
            home = session.get(BLS_HOME, timeout=60)
            home.raise_for_status()
            token_match = re.search(r"[?&]token=([A-Za-z0-9-]{8,})", home.text)
        if not token_match:
            raise RuntimeError("Could not discover the official BLS download token")
        zip_url = f"https://blsdb.de/assets/uploads/BLS_4_0_2025_DE.zip?token={token_match.group(1)}"
    else:
        from urllib.parse import urljoin
        zip_url = urljoin(BLS_HOME, match.group(1).replace("&amp;", "&"))

    response = session.get(zip_url, timeout=180)
    response.raise_for_status()
    archive = zipfile.ZipFile(io.BytesIO(response.content))
    candidates = [
        name
        for name in archive.namelist()
        if name.lower().endswith(".xlsx")
        and "daten" in name.lower()
        and "mapping" not in name.lower()
        and "components" not in name.lower()
    ]
    if not candidates:
        raise RuntimeError("Official BLS ZIP did not contain the expected data workbook")
    return archive.read(candidates[0])


def read_canonical_codes(path: Path | None) -> tuple[set[str], dict[str, str]]:
    if path:
        raw = path.read_text(encoding="utf-8")
    else:
        response = requests.get(CANONICAL_ITEMS_URL, timeout=60)
        response.raise_for_status()
        raw = response.text

    selected: set[str] = set()
    names: dict[str, str] = {}
    reader = csv.DictReader(io.StringIO(raw))
    for row in reader:
        code = (row.get("code") or "").strip()
        if not code or code[0] not in ALLOWED_MAIN_GROUPS:
            continue
        selected.add(code)
        names[code] = (row.get("name_de") or "").strip()

    if len(selected) != EXPECTED_FOODS:
        raise RuntimeError(
            f"Canonical selection changed: expected {EXPECTED_FOODS} B-W codes, got {len(selected)}"
        )
    return selected, names


def value_as_text(value: Any) -> str | None:
    """Preserve BLS cell semantics without coercing missing/special tokens to zero."""
    if value is None:
        return None
    if isinstance(value, str):
        return value
    # openpyxl returns Excel numerics as int/float. str/repr changes formatting only,
    # never the numeric value; no rounding or unit conversion is performed.
    if isinstance(value, float):
        return repr(value)
    return str(value)


def parse_nutrient_header(header: str) -> tuple[str, str, str | None]:
    header = str(header or "").strip()
    if not header:
        raise RuntimeError("Empty nutrient header")
    code, _, remainder = header.partition(" ")
    remainder = remainder.strip()
    unit = None
    match = re.search(r"\s*\[([^\]]+)\]\s*$", remainder)
    if match:
        unit_raw = match.group(1).strip()
        unit = re.sub(r"/\s*100\s*g\s*$", "", unit_raw, flags=re.IGNORECASE).strip()
        name = remainder[: match.start()].strip()
    else:
        name = remainder
    return code, name, unit


def build_database(xlsx_bytes: bytes, selected: set[str], output: Path) -> dict[str, Any]:
    workbook = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    header = next(rows)

    if tuple(header[:3]) != ("BLS Code", "Lebensmittelbezeichnung", "Food name"):
        raise RuntimeError(f"Unexpected BLS workbook header: {header[:3]!r}")
    if (len(header) - 3) % 3 != 0:
        raise RuntimeError(f"Unexpected BLS column count: {len(header)}")

    nutrient_columns: list[tuple[int, str, str, str | None, str]] = []
    for col in range(3, len(header), 3):
        source_header = str(header[col] or "").strip()
        nutrient_code, nutrient_name, unit = parse_nutrient_header(source_header)
        nutrient_columns.append((col, nutrient_code, nutrient_name, unit, source_header))

    if len(nutrient_columns) != EXPECTED_NUTRIENTS:
        raise RuntimeError(
            f"Expected {EXPECTED_NUTRIENTS} nutrient columns, got {len(nutrient_columns)}"
        )
    codes = [x[1] for x in nutrient_columns]
    if len(codes) != len(set(codes)):
        duplicates = [k for k, v in Counter(codes).items() if v > 1]
        raise RuntimeError(f"Duplicate nutrient component codes: {duplicates}")

    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists():
        output.unlink()

    db = sqlite3.connect(output)
    db.execute("PRAGMA journal_mode=DELETE")
    db.execute("PRAGMA foreign_keys=ON")
    db.execute("PRAGMA synchronous=FULL")
    db.executescript(
        """
        CREATE TABLE meta (
            key TEXT PRIMARY KEY,
            value TEXT NOT NULL
        ) WITHOUT ROWID;

        CREATE TABLE foods (
            bls_code TEXT PRIMARY KEY,
            name_de TEXT NOT NULL
        ) WITHOUT ROWID;

        CREATE TABLE nutrients (
            ordinal INTEGER PRIMARY KEY,
            code TEXT NOT NULL UNIQUE,
            name_de TEXT NOT NULL,
            unit TEXT,
            source_header TEXT NOT NULL
        );

        CREATE TABLE food_nutrients (
            bls_code TEXT NOT NULL,
            nutrient_code TEXT NOT NULL,
            value_raw TEXT,
            PRIMARY KEY (bls_code, nutrient_code),
            FOREIGN KEY (bls_code) REFERENCES foods(bls_code),
            FOREIGN KEY (nutrient_code) REFERENCES nutrients(code)
        ) WITHOUT ROWID;

        CREATE INDEX idx_foods_name_de ON foods(name_de COLLATE NOCASE);
        CREATE INDEX idx_food_nutrients_nutrient ON food_nutrients(nutrient_code);
        """
    )

    db.executemany(
        "INSERT INTO nutrients(ordinal, code, name_de, unit, source_header) VALUES (?, ?, ?, ?, ?)",
        [
            (ordinal, code, name, unit, source_header)
            for ordinal, (_, code, name, unit, source_header) in enumerate(nutrient_columns, start=1)
        ],
    )

    found: set[str] = set()
    special_counter: Counter[str] = Counter()
    null_cells = 0
    total_cells = 0

    for row in rows:
        if not row or row[0] is None:
            continue
        bls_code = str(row[0]).strip()
        if bls_code not in selected:
            continue
        if bls_code in found:
            raise RuntimeError(f"Duplicate selected BLS code in workbook: {bls_code}")

        name_de = str(row[1] or "").strip()
        if not name_de:
            raise RuntimeError(f"Selected BLS row has no German name: {bls_code}")
        db.execute("INSERT INTO foods(bls_code, name_de) VALUES (?, ?)", (bls_code, name_de))

        nutrient_rows = []
        for col, nutrient_code, _, _, _ in nutrient_columns:
            value_raw = value_as_text(row[col] if col < len(row) else None)
            total_cells += 1
            if value_raw is None or value_raw == "":
                null_cells += 1
            elif not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", value_raw):
                special_counter[value_raw] += 1
            nutrient_rows.append((bls_code, nutrient_code, value_raw))

        db.executemany(
            "INSERT INTO food_nutrients(bls_code, nutrient_code, value_raw) VALUES (?, ?, ?)",
            nutrient_rows,
        )
        found.add(bls_code)

    missing = selected - found
    if missing:
        raise RuntimeError(f"{len(missing)} selected BLS codes missing from workbook: {sorted(missing)[:20]}")
    if len(found) != EXPECTED_FOODS:
        raise RuntimeError(f"Expected {EXPECTED_FOODS} foods in database, got {len(found)}")

    meta = {
        "source": "Bundeslebensmittelschlüssel (BLS) 4.0 (2025)",
        "publisher": "Max Rubner-Institut",
        "license": "CC BY 4.0",
        "attribution": ATTRIBUTION,
        "doi": "10.25826/Data20251217-134202-0",
        "basis": "all nutrient values per 100 g edible portion",
        "selection": "EiSiMo/bls-icons canonical items.csv; BLS main groups B-W only; X/Y menu components excluded",
        "foods": str(len(found)),
        "nutrients": str(len(nutrient_columns)),
        "value_storage": "value_raw TEXT; no rounding, no unit conversion, no TR/<LOD/<LOQ/- coercion",
    }
    db.executemany("INSERT INTO meta(key, value) VALUES (?, ?)", meta.items())
    db.commit()

    integrity = db.execute("PRAGMA integrity_check").fetchone()[0]
    row_count = db.execute("SELECT COUNT(*) FROM foods").fetchone()[0]
    nutrient_count = db.execute("SELECT COUNT(*) FROM nutrients").fetchone()[0]
    value_count = db.execute("SELECT COUNT(*) FROM food_nutrients").fetchone()[0]

    samples = {}
    for code in ("F503100", "M012200", "M402600"):
        row = db.execute("SELECT name_de FROM foods WHERE bls_code=?", (code,)).fetchone()
        if row:
            samples[code] = row[0]

    db.close()

    expected_values = EXPECTED_FOODS * EXPECTED_NUTRIENTS
    if integrity != "ok" or row_count != EXPECTED_FOODS or nutrient_count != EXPECTED_NUTRIENTS:
        raise RuntimeError(
            f"Validation failed: integrity={integrity}, foods={row_count}, nutrients={nutrient_count}"
        )
    if value_count != expected_values:
        raise RuntimeError(f"Expected {expected_values} food-nutrient cells, got {value_count}")

    return {
        "integrity_check": integrity,
        "foods": row_count,
        "nutrients": nutrient_count,
        "food_nutrient_cells": value_count,
        "null_cells": null_cells,
        "special_tokens": dict(sorted(special_counter.items())),
        "sample_foods": samples,
        "database_bytes": output.stat().st_size,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--bls-xlsx", type=Path, help="Optional local BLS_4_0_Daten_2025_DE.xlsx")
    parser.add_argument("--canonical-items", type=Path, help="Optional local EiSiMo/bls-icons items.csv")
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("BLS_4_0_ingredients_2992.sqlite"),
    )
    parser.add_argument("--report", type=Path, default=Path("BLS_4_0_ingredients_2992.validation.json"))
    args = parser.parse_args()

    selected, _ = read_canonical_codes(args.canonical_items)
    xlsx_bytes = args.bls_xlsx.read_bytes() if args.bls_xlsx else download_official_xlsx()
    report = build_database(xlsx_bytes, selected, args.output)
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
